import os
import json
import re
import qrcode
import tempfile
import shutil
from openpyxl import Workbook
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

# Paths
album_dir = r'C:\Users\Admin\Desktop\Album'
alumnos_js = os.path.join(album_dir, 'data', 'alumnos.js')
output_excel = os.path.join(album_dir, 'QRs_Alumnos.xlsx')

# Read JS JSON data
with open(alumnos_js, 'r', encoding='utf-8') as f:
    content = f.read()
start_idx = content.find('{')
end_idx = content.rfind('}') + 1
json_str = content[start_idx:end_idx]
data = json.loads(json_str)

# Filter students who have valid local images (starts with 'img/')
alumnos_validos = []
for a in data['alumnos']:
    if a.get('img') and a['img'].startswith('img/'):
        alumnos_validos.append(a)

print(f"Total alumnos con imagen: {len(alumnos_validos)}")

# Group by 'grupo'
alumnos_por_grupo = {}
for a in alumnos_validos:
    grupo = a['grupo']
    if grupo not in alumnos_por_grupo:
        alumnos_por_grupo[grupo] = []
    alumnos_por_grupo[grupo].append(a)

# Create Excel Workbook
wb = Workbook()
# Remove default sheet
default_sheet = wb.active
wb.remove(default_sheet)

# Design Excel style
font_header = Font(name="Arial", size=11, bold=True, color="FFFFFF")
font_body = Font(name="Arial", size=10)
fill_header = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
align_left = Alignment(horizontal="left", vertical="center")

thin_side = Side(border_style="thin", color="D9D9D9")
border_cell = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

headers = [
    "ID Alumno", 
    "Nombre", 
    "Grupo", 
    "Ruta Foto", 
    "URL Google Drive", 
    "Código QR (ID)", 
    "Código QR (Drive)"
]

# Temporary directory for QR images
temp_dir = tempfile.mkdtemp()

# Create a sheet for each group
for grupo, alumnos in sorted(alumnos_por_grupo.items()):
    # Clean sheet name (Excel limits to 31 chars and bans certain symbols)
    sheet_title = re.sub(r'[\\*?:/\[\]]', '', grupo)[:31]
    
    ws = wb.create_sheet(title=sheet_title)
    
    # Headers
    ws.append(headers)
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center

    # Column widths
    column_widths = {
        'A': 15,  # ID
        'B': 25,  # Nombre
        'C': 15,  # Grupo
        'D': 25,  # Ruta Foto
        'E': 30,  # URL Drive
        'F': 20,  # QR ID
        'G': 20   # QR Drive
    }
    for col, width in column_widths.items():
        ws.column_dimensions[col].width = width

    row_idx = 2
    for a in alumnos:
        student_id = a['id']
        nombre = a['nombre']
        img_path = a['img']
        drive_url = a.get('drive_url', '')

        # Row height
        ws.row_dimensions[row_idx].height = 110

        # Write text
        ws.cell(row=row_idx, column=1, value=student_id).alignment = align_center
        ws.cell(row=row_idx, column=2, value=nombre).alignment = align_left
        ws.cell(row=row_idx, column=3, value=grupo).alignment = align_center
        ws.cell(row=row_idx, column=4, value=img_path).alignment = align_left
        ws.cell(row=row_idx, column=5, value=drive_url).alignment = align_left

        # Borders
        for col_num in range(1, 8):
            cell = ws.cell(row=row_idx, column=col_num)
            cell.font = font_body
            cell.border = border_cell

        # QR ID
        qr_id_path = os.path.join(temp_dir, f"qr_id_{student_id}.png")
        qr = qrcode.QRCode(version=1, box_size=3, border=1)
        qr.add_data(student_id)
        qr.make(fit=True)
        img_qr_id = qr.make_image(fill_color="black", back_color="white")
        img_qr_id.save(qr_id_path)

        xl_img_id = OpenpyxlImage(qr_id_path)
        ws.add_image(xl_img_id, f"F{row_idx}")

        # QR Drive
        if drive_url:
            qr_drive_path = os.path.join(temp_dir, f"qr_drive_{student_id}.png")
            qr_d = qrcode.QRCode(version=1, box_size=3, border=1)
            qr_d.add_data(drive_url)
            qr_d.make(fit=True)
            img_qr_drive = qr_d.make_image(fill_color="black", back_color="white")
            img_qr_drive.save(qr_drive_path)

            xl_img_drive = OpenpyxlImage(qr_drive_path)
            ws.add_image(xl_img_drive, f"G{row_idx}")
        else:
            ws.cell(row=row_idx, column=7, value="Sin URL").alignment = align_center

        row_idx += 1

# Save Workbook
wb.save(output_excel)

# Clean up
try:
    shutil.rmtree(temp_dir)
except Exception as e:
    print(f"Error deleting temp dir: {e}")

print(f"Excel file generated successfully with group sheets: {output_excel}")
